"""Build the editable resume; optionally render it with the Codex DOCX renderer.

Run with the bundled document Python runtime (python-docx). To produce the PDF,
pass --render-script with the absolute path to the bundled render_docx.py. That
renderer selects bundled LibreOffice and Poppler, never desktop LibreOffice.
Render intermediates stay outside static/downloads. Inspect every rendered page
after changing content or layout.
"""

from __future__ import annotations

import argparse
from pathlib import Path
import shutil
import subprocess
import sys

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
STEM = "saiganesh-sattenapalli-resume"
GITHUB = "https://github.com/saiganeshsattenapalli"
PORTFOLIO_SOURCE = GITHUB + "/My-Portfolio"


def hyperlink(paragraph, text: str, url: str, *, size: float = 10.5,
              bold: bool = False, color: str = "1E4E6F") -> None:
    relationship = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), relationship)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Arial")
    fonts.set(qn("w:hAnsi"), "Arial")
    properties.append(fonts)
    size_element = OxmlElement("w:sz")
    size_element.set(qn("w:val"), str(int(size * 2)))
    properties.append(size_element)
    color_element = OxmlElement("w:color")
    color_element.set(qn("w:val"), color)
    properties.append(color_element)
    if bold:
        properties.append(OxmlElement("w:b"))
    run.append(properties)
    content = OxmlElement("w:t")
    content.text = text
    run.append(content)
    link.append(run)
    paragraph._p.append(link)


def plain_style(style, size: float, *, bold: bool = False) -> None:
    style.font.name = "Arial"
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.font.underline = False
    style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), "Arial")


def add_bullet(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(text, style="List Bullet")
    paragraph.paragraph_format.keep_together = True


def add_project(document: Document, name: str, years: str, url: str,
                descriptor: str, bullets: list[str]) -> None:
    paragraph = document.add_paragraph(style="Project")
    paragraph.add_run(name).bold = True
    paragraph.add_run("  ")
    hyperlink(paragraph, "GitHub", url, size=9.5)
    paragraph.add_run("\t" + years)
    paragraph = document.add_paragraph(descriptor, style="Project Detail")
    for bullet in bullets:
        add_bullet(document, bullet)


def build(output_dir: Path) -> Path:
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.53)
    section.bottom_margin = Inches(0.53)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)

    normal = document.styles["Normal"]
    plain_style(normal, 11)
    normal.paragraph_format.space_after = Pt(3)
    normal.paragraph_format.line_spacing = 1.08
    normal.paragraph_format.widow_control = True

    title = document.styles["Title"]
    plain_style(title, 24, bold=True)
    title.paragraph_format.space_after = Pt(3)
    title.paragraph_format.line_spacing = 1
    title.paragraph_format.keep_with_next = True

    subtitle = document.styles["Subtitle"]
    plain_style(subtitle, 11.5)
    subtitle.paragraph_format.space_after = Pt(6)
    subtitle.paragraph_format.keep_with_next = True

    heading = document.styles["Heading 1"]
    plain_style(heading, 11, bold=True)
    heading.paragraph_format.space_before = Pt(10)
    heading.paragraph_format.space_after = Pt(5)
    heading.paragraph_format.keep_with_next = True

    bullet = document.styles["List Bullet"]
    plain_style(bullet, 11)
    bullet.paragraph_format.left_indent = Inches(0.13)
    bullet.paragraph_format.first_line_indent = Inches(-0.13)
    bullet.paragraph_format.space_after = Pt(3.5)
    bullet.paragraph_format.line_spacing = 1.07

    for name in ("Project", "Project Detail", "Contact", "Skill"):
        document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        document.styles[name].base_style = normal
    project = document.styles["Project"]
    plain_style(project, 11)
    project.paragraph_format.space_before = Pt(5)
    project.paragraph_format.space_after = Pt(2)
    project.paragraph_format.keep_with_next = True
    project.paragraph_format.tab_stops.add_tab_stop(Inches(7.2), WD_TAB_ALIGNMENT.RIGHT)
    detail = document.styles["Project Detail"]
    plain_style(detail, 9.5)
    detail.paragraph_format.space_after = Pt(4)
    detail.paragraph_format.keep_with_next = True
    contact = document.styles["Contact"]
    plain_style(contact, 10)
    contact.paragraph_format.space_after = Pt(3)
    contact.paragraph_format.keep_with_next = True
    document.styles["Skill"].paragraph_format.space_after = Pt(3)

    # Clear theme formatting and border residue from title/section styles.
    for style in (title, subtitle, heading):
        for element in list(style._element.iter()):
            if element.tag in (qn("w:pBdr"), qn("w:shd")):
                element.getparent().remove(element)
            for attribute in ("themeColor", "themeTint", "themeShade"):
                element.attrib.pop(qn("w:" + attribute), None)

    document.add_paragraph("Saiganesh Sattenapalli", "Title")
    document.add_paragraph("Data Analysis and AI Engineering", "Subtitle")
    paragraph = document.add_paragraph(style="Contact")
    paragraph.add_run("Hyderabad, India  |  ")
    hyperlink(paragraph, "+91 9346753626", "tel:+919346753626", size=10)
    paragraph.add_run("  |  ")
    hyperlink(paragraph, "saiganeshsattenapalli26@gmail.com",
              "mailto:saiganeshsattenapalli26@gmail.com", size=10)
    paragraph = document.add_paragraph(style="Contact")
    hyperlink(paragraph, "github.com/saiganeshsattenapalli", GITHUB, size=10)
    paragraph.add_run("  |  ")
    hyperlink(paragraph, "Portfolio source", PORTFOLIO_SOURCE, size=10)

    document.add_paragraph("Profile", "Heading 1")
    document.add_paragraph(
        "Second-year computer science student building the Curiora ecosystem of local AI "
        "applications and data workflows with Python and SQL. Project work spans multimodal model integration, "
        "transaction reporting, and classification prototypes. Seeking Data Analyst, "
        "ML Engineer, and AI Engineer opportunities."
    )

    document.add_paragraph("Education", "Heading 1")
    paragraph = document.add_paragraph()
    paragraph.add_run("B.Tech in Computer Science (AIML)").bold = True
    paragraph.add_run("  |  2025-2029 (expected)")
    document.add_paragraph("Malla Reddy College of Engineering, Hyderabad")

    document.add_paragraph("Technical Skills", "Heading 1")
    for label, skills in [
        ("Data and programming", "Python, SQL, pandas, NumPy, PostgreSQL"),
        ("Machine learning", "scikit-learn, logistic regression, feature preparation"),
        ("AI applications", "MLX, PyTorch, Hugging Face Transformers, text and vision model integration"),
        ("Application development", "FastAPI, Git, Jinja2, JavaScript, HTML, CSS"),
    ]:
        paragraph = document.add_paragraph(style="Skill")
        paragraph.add_run(label + ": ").bold = True
        paragraph.add_run(skills)

    document.add_paragraph("Selected Projects", "Heading 1")
    add_project(
        document, "Curiora CampusOS", "2026", GITHUB + "/Curiora-CampusOS",
        "Local AI application prototype | Python, FastAPI, MLX, PyTorch",
        [
            "Built a unified text and image inference API integrating GPT-OSS-20B and "
            "Qwen3-VL-8B, with runtime selection for Apple Silicon MLX, CUDA, and CPU.",
            "Implemented model unloading before model switches, conversation IDs, image "
            "validation, temporary file cleanup, and inference dispatched to a thread pool.",
        ],
    )
    add_project(
        document, "Curiora Pay", "2025-2026",
        "https://github.com/Curiora-intelligence/Curiora-Pay",
        "Banking and data workflow prototype | Python, PostgreSQL, pandas, scikit-learn",
        [
            "Implemented deposits, withdrawals, transfers, and transaction history; "
            "queried account transactions with SQL and exported statements through pandas to CSV.",
            "Integrated a logistic regression demonstration using balance and transaction "
            "count as features, trained on a small synthetic dataset for learning purposes.",
        ],
    )
    add_project(
        document, "Curiora Research", "2026", GITHUB + "/Curiora-Research",
        "Local visual intelligence experiment | Python, MLX, Qwen3-VL",
        [
            "Implemented lazy loading and serialized inference for a local vision-language "
            "model, with prompts that require image-grounded answers and explicit uncertainty.",
        ],
    )

    document.core_properties.title = "Saiganesh Sattenapalli Resume"
    document.core_properties.subject = "Data Analyst ML Engineer and AI Engineer roles"
    document.core_properties.author = "Saiganesh Sattenapalli"
    document.core_properties.keywords = "Python, SQL, Machine Learning, Data Analysis, AI, FastAPI"
    document.core_properties.comments = ""
    output_dir.mkdir(parents=True, exist_ok=True)
    path = output_dir / (STEM + ".docx")
    document.save(path)
    return path


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output-dir", type=Path, default=ROOT / "static" / "downloads")
    parser.add_argument("--render-script", type=Path)
    parser.add_argument("--qa-dir", type=Path, default=Path("/private/tmp/portfolio-resume/render"))
    args = parser.parse_args()
    docx = build(args.output_dir)
    print(docx)
    if args.render_script:
        if not args.render_script.is_file():
            parser.error("--render-script must point to the bundled render_docx.py")
        subprocess.run([
            sys.executable, str(args.render_script), str(docx),
            "--output_dir", str(args.qa_dir), "--emit_pdf",
        ], check=True)
        pdf = args.qa_dir / (STEM + ".pdf")
        destination = args.output_dir / pdf.name
        shutil.copyfile(pdf, destination)
        print(destination)


if __name__ == "__main__":
    main()
